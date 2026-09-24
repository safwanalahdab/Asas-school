from datetime import date
from io import BytesIO
import os
from pathlib import Path
from unittest import skipUnless
import zipfile

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl import Workbook

from students.models import (
    Enrollment,
    GuardianStudent,
    Student,
    StudentHealthProfile,
    StudentImportJob,
    StudentImportRow,
)
from students.student_import_excel import (
    EXPECTED_HEADERS,
    IMPORT_SHEET_NAME,
    MAX_DATA_ROWS,
    MAX_FILE_SIZE_BYTES,
    MAX_XLSX_PART_UNCOMPRESSED_BYTES,
    MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES,
    MAX_XLSX_ZIP_ENTRIES,
    StudentImportReadResult,
    _validate_xlsx_archive,
    read_student_import_xlsx,
)


REAL_TEMPLATE_PATH = Path(
    os.environ.get(
        "STUDENT_IMPORT_TEMPLATE_PATH",
        r"E:\Safwan\Downloads\قالب_استيراد_طلاب_مدرسة_أساس (1).xlsx",
    )
)
User = get_user_model()


def valid_row(**overrides):
    values = {
        "الاسم الأول بالعربية *": "رامي",
        "الكنية بالعربية *": "اختبار",
        "الاسم الأول بالإنكليزية": "Rami",
        "الكنية بالإنكليزية": "Test",
        "اسم الأب": "ماهر",
        "اسم الأم": "هند",
        "تاريخ الميلاد *": date(2018, 5, 10),
        "الجنس *": "ذكر",
        "الرقم الوطني لولي الأمر": "٠١٢٣٤٥٦٧٨٩١",
        "اسم ولي الأمر الأول": "ماهر",
        "كنية ولي الأمر": "اختبار",
        "رقم هاتف ولي الأمر": "۰۹۳۳۱۱۱۲۲۲",
        "صلة القرابة": "أب",
        "زمرة الدم": "O+",
        "الأمراض المزمنة": "",
        "الحساسية": "",
        "الأدوية الدائمة": "",
        "الاحتياجات الصحية الخاصة": "",
        "اسم جهة اتصال للطوارئ": "ماهر اختبار",
        "هاتف الطوارئ": "0944000000",
        "ملاحظات صحية": "",
        "السنة الدراسية *": "٢٠٢٦/٢٠٢٧",
        "المرحلة *": "الابتدائية",
        "الصف *": "الصف الثالث",
        "الشعبة *": "أ",
        "تاريخ التسجيل *": "2026-09-01",
        "طريقة الحضور المعتادة": "",
        "طريقة الانصراف المعتادة": "ولي الأمر",
    }
    values.update(overrides)
    return [values[header] for header in EXPECTED_HEADERS]


def workbook_bytes(*rows, sheet_name=IMPORT_SHEET_NAME, headers=EXPECTED_HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


def error_codes(errors):
    return {error.code for error in errors}


def zip_entry(name, *, size=0, compressed_size=None):
    entry = zipfile.ZipInfo(name)
    entry.file_size = size
    entry.compress_size = size if compressed_size is None else compressed_size
    return entry


class ArchiveMetadata:
    def __init__(self, entries):
        self.entries = entries

    def infolist(self):
        return self.entries


class StudentImportExcelReaderTests(TestCase):
    def read(self, content, filename="students.xlsx"):
        return read_student_import_xlsx(content, filename=filename)

    def test_valid_file_with_one_row_is_normalized(self):
        result = self.read(workbook_bytes(valid_row()))

        self.assertIsInstance(result, StudentImportReadResult)
        self.assertTrue(result.is_valid)
        self.assertEqual(result.data_row_count, 1)
        self.assertEqual(result.rows[0].excel_row_number, 2)
        self.assertEqual(result.rows[0].normalized_data["student"]["gender"], "male")
        self.assertEqual(result.rows[0].normalized_data["guardian"]["national_id"], "01234567891")
        self.assertEqual(result.rows[0].normalized_data["guardian"]["phone_number"], "0933111222")
        self.assertEqual(result.rows[0].normalized_data["enrollment"]["academic_year"], "2026/2027")
        self.assertEqual(result.rows[0].normalized_data["enrollment"]["usual_arrival_method"], "school_bus")
        self.assertEqual(result.rows[0].normalized_data["enrollment"]["usual_departure_method"], "guardian")

    def test_valid_file_with_multiple_rows_preserves_excel_row_numbers(self):
        result = self.read(workbook_bytes(valid_row(), [None] * 28, valid_row()))

        self.assertTrue(result.is_valid)
        self.assertEqual(result.data_row_count, 2)
        self.assertEqual([row.excel_row_number for row in result.rows], [2, 4])

    def test_wrong_sheet_name_is_rejected(self):
        result = self.read(workbook_bytes(valid_row(), sheet_name="طلاب"))

        self.assertEqual(error_codes(result.file_errors), {"missing_sheet"})

    def test_missing_or_changed_header_is_rejected(self):
        changed_headers = list(EXPECTED_HEADERS)
        changed_headers[3] = "عنوان مختلف"

        result = self.read(workbook_bytes(valid_row(), headers=changed_headers))

        self.assertEqual(error_codes(result.file_errors), {"invalid_headers"})

    def test_corrupt_or_unsupported_file_is_rejected(self):
        corrupt = self.read(BytesIO(b"not an xlsx"))
        corrupt_zip = self.read(BytesIO(b"PK\x03\x04" + (b"\x00" * 30)))
        unsupported = self.read(workbook_bytes(valid_row()), filename="students.xls")

        self.assertEqual(error_codes(corrupt.file_errors), {"invalid_xlsx"})
        self.assertEqual(error_codes(corrupt_zip.file_errors), {"invalid_xlsx"})
        self.assertEqual(error_codes(unsupported.file_errors), {"unsupported_extension"})

    def test_oversized_file_is_rejected_before_parsing(self):
        oversized = BytesIO(b"x" * (MAX_FILE_SIZE_BYTES + 1))

        result = self.read(oversized)

        self.assertEqual(error_codes(result.file_errors), {"file_too_large"})

    def test_zip_with_too_many_entries_is_rejected_from_metadata(self):
        archive = ArchiveMetadata(
            [zip_entry(f"part-{index}.xml") for index in range(MAX_XLSX_ZIP_ENTRIES + 1)]
        )

        error = _validate_xlsx_archive(archive)

        self.assertEqual(error.code, "unsafe_xlsx_archive")

    def test_zip_with_oversized_uncompressed_part_is_rejected_from_metadata(self):
        archive = ArchiveMetadata(
            [
                zip_entry(
                    "xl/worksheets/sheet1.xml",
                    size=MAX_XLSX_PART_UNCOMPRESSED_BYTES + 1,
                    compressed_size=1,
                )
            ]
        )

        error = _validate_xlsx_archive(archive)

        self.assertEqual(error.code, "unsafe_xlsx_archive")

    def test_zip_with_oversized_uncompressed_total_is_rejected_from_metadata(self):
        part_size = MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES // 3
        archive = ArchiveMetadata(
            [
                zip_entry(f"part-{index}.xml", size=part_size, compressed_size=part_size)
                for index in range(4)
            ]
        )

        error = _validate_xlsx_archive(archive)

        self.assertEqual(error.code, "unsafe_xlsx_archive")

    def test_empty_file_and_workbook_without_data_are_rejected(self):
        empty_file = self.read(BytesIO())
        no_rows = self.read(workbook_bytes())

        self.assertEqual(error_codes(empty_file.file_errors), {"empty_file"})
        self.assertEqual(error_codes(no_rows.file_errors), {"no_data_rows"})

    def test_exactly_1500_rows_are_allowed(self):
        result = self.read(workbook_bytes(*(valid_row() for _ in range(MAX_DATA_ROWS))))

        self.assertFalse(result.file_errors)
        self.assertEqual(result.data_row_count, MAX_DATA_ROWS)

    def test_data_after_1500_rows_is_rejected(self):
        result = self.read(
            workbook_bytes(*(valid_row() for _ in range(MAX_DATA_ROWS + 1)))
        )

        self.assertEqual(
            error_codes(result.file_errors),
            {"data_outside_template_range"},
        )

    def test_trailing_formatted_empty_rows_are_not_counted(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = IMPORT_SHEET_NAME
        sheet.append(list(EXPECTED_HEADERS))
        sheet.append(valid_row())
        sheet.cell(row=2000, column=1).number_format = "@"
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        result = self.read(output)

        self.assertFalse(result.file_errors)
        self.assertEqual(result.data_row_count, 1)

    def test_field_errors_use_actual_excel_row_number(self):
        invalid = valid_row(**{"الجنس *": "غير معروف"})
        result = self.read(workbook_bytes([None] * 28, invalid))

        error = result.rows[0].validation_errors[0]
        self.assertEqual(result.rows[0].excel_row_number, 3)
        self.assertEqual(error.excel_row_number, 3)
        self.assertEqual(error.field, "gender")
        self.assertEqual(error.column, "الجنس *")

    def test_missing_required_student_field_is_reported(self):
        result = self.read(workbook_bytes(valid_row(**{"الاسم الأول بالعربية *": ""})))

        self.assertIn("required", error_codes(result.rows[0].validation_errors))
        self.assertEqual(result.rows[0].validation_errors[0].field, "first_name")

    def test_partial_guardian_requires_identity_name_and_relationship(self):
        row = valid_row(
            **{
                "الرقم الوطني لولي الأمر": "",
                "اسم ولي الأمر الأول": "ماهر",
                "كنية ولي الأمر": "",
                "رقم هاتف ولي الأمر": "",
                "صلة القرابة": "",
            }
        )

        result = self.read(workbook_bytes(row))
        error_fields = {error.field for error in result.rows[0].validation_errors}

        self.assertEqual(
            error_fields,
            {"guardian_national_id", "guardian_last_name", "guardian_relationship"},
        )

    def test_student_without_guardian_is_allowed(self):
        row = valid_row(
            **{
                "الرقم الوطني لولي الأمر": "",
                "اسم ولي الأمر الأول": "",
                "كنية ولي الأمر": "",
                "رقم هاتف ولي الأمر": "",
                "صلة القرابة": "",
            }
        )

        result = self.read(workbook_bytes(row))

        self.assertTrue(result.is_valid)
        self.assertIsNone(result.rows[0].normalized_data["guardian"])

    def test_invalid_gender_blood_type_and_transportation_are_reported(self):
        row = valid_row(
            **{
                "الجنس *": "آخر",
                "زمرة الدم": "X+",
                "طريقة الحضور المعتادة": "سيارة",
            }
        )

        result = self.read(workbook_bytes(row))

        self.assertEqual(
            error_codes(result.rows[0].validation_errors),
            {"invalid_choice", "invalid_blood_type"},
        )

    def test_valid_date_and_ambiguous_or_invalid_date(self):
        valid = self.read(
            workbook_bytes(valid_row(**{"تاريخ الميلاد *": "2018-05-10"}))
        )
        ambiguous = self.read(
            workbook_bytes(valid_row(**{"تاريخ الميلاد *": "10/05/2018"}))
        )

        self.assertEqual(
            valid.rows[0].normalized_data["student"]["birth_date"],
            "2018-05-10",
        )
        self.assertIn("invalid_date", error_codes(ambiguous.rows[0].validation_errors))

    def test_formula_in_student_data_is_rejected(self):
        result = self.read(
            workbook_bytes(valid_row(**{"الاسم الأول بالعربية *": "=1+1"}))
        )

        self.assertIn("formula_not_allowed", error_codes(result.rows[0].validation_errors))

    def test_reader_does_not_create_school_data(self):
        result = self.read(workbook_bytes(valid_row()))

        self.assertTrue(result.is_valid)
        self.assertEqual(User.objects.count(), 0)
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(GuardianStudent.objects.count(), 0)
        self.assertEqual(StudentHealthProfile.objects.count(), 0)
        self.assertEqual(Enrollment.objects.count(), 0)
        self.assertEqual(StudentImportJob.objects.count(), 0)
        self.assertEqual(StudentImportRow.objects.count(), 0)

    @skipUnless(REAL_TEMPLATE_PATH.exists(), "Real student import template is unavailable.")
    def test_real_template_has_the_approved_structure(self):
        result = read_student_import_xlsx(REAL_TEMPLATE_PATH)

        self.assertNotIn("missing_sheet", error_codes(result.file_errors))
        self.assertNotIn("invalid_headers", error_codes(result.file_errors))
