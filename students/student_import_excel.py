from dataclasses import asdict, dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import posixpath
import re
import zipfile
from xml.etree.ElementTree import ParseError, fromstring, iterparse

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


IMPORT_SHEET_NAME = "إدخال الطلاب"
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 1500
MAX_TEMPLATE_EXCEL_ROW = MAX_DATA_ROWS + 1
MAX_XLSX_ZIP_ENTRIES = 1000
MAX_XLSX_PART_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 1000

SPREADSHEET_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
PACKAGE_RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)

EXPECTED_HEADERS = (
    "الاسم الأول بالعربية *",
    "الكنية بالعربية *",
    "الاسم الأول بالإنكليزية",
    "الكنية بالإنكليزية",
    "اسم الأب",
    "اسم الأم",
    "تاريخ الميلاد *",
    "الجنس *",
    "الرقم الوطني لولي الأمر",
    "اسم ولي الأمر الأول",
    "كنية ولي الأمر",
    "رقم هاتف ولي الأمر",
    "صلة القرابة",
    "زمرة الدم",
    "الأمراض المزمنة",
    "الحساسية",
    "الأدوية الدائمة",
    "الاحتياجات الصحية الخاصة",
    "اسم جهة اتصال للطوارئ",
    "هاتف الطوارئ",
    "ملاحظات صحية",
    "السنة الدراسية *",
    "المرحلة *",
    "الصف *",
    "الشعبة *",
    "تاريخ التسجيل *",
    "طريقة الحضور المعتادة",
    "طريقة الانصراف المعتادة",
)

FIELD_NAMES = (
    "first_name",
    "last_name",
    "first_name_en",
    "last_name_en",
    "father_name",
    "mother_name",
    "birth_date",
    "gender",
    "guardian_national_id",
    "guardian_first_name",
    "guardian_last_name",
    "guardian_phone_number",
    "guardian_relationship",
    "blood_type",
    "chronic_diseases",
    "allergies",
    "permanent_medications",
    "special_health_needs",
    "emergency_contact_name",
    "emergency_contact_phone",
    "health_notes",
    "academic_year",
    "stage",
    "grade_level",
    "section",
    "enrollment_date",
    "usual_arrival_method",
    "usual_departure_method",
)

HEADER_BY_FIELD = dict(zip(FIELD_NAMES, EXPECTED_HEADERS))
ARABIC_DIGITS_TRANSLATION = str.maketrans(
    "٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹",
    "01234567890123456789",
)
GENDER_VALUES = {"ذكر": "male", "أنثى": "female"}
STAGE_VALUES = {
    "الروضة": "kindergarten",
    "الابتدائية": "primary",
    "الإعدادية": "preparatory",
    "الثانوية": "secondary",
}
TRANSPORTATION_VALUES = {
    "باص المدرسة": "school_bus",
    "ولي الأمر": "guardian",
}
BLOOD_TYPES = {"A+", "A-", "B+", "B-", "AB+", "AB-", "O+", "O-"}
ACADEMIC_YEAR_PATTERN = re.compile(r"^(\d{4})/(\d{4})$")


@dataclass(frozen=True)
class ImportValidationError:
    code: str
    detail: str
    field: str | None = None
    column: str | None = None
    excel_row_number: int | None = None

    def to_dict(self):
        return {key: value for key, value in asdict(self).items() if value is not None}


@dataclass(frozen=True)
class StudentImportRowResult:
    excel_row_number: int
    normalized_data: dict
    validation_errors: list[ImportValidationError]

    def to_dict(self):
        return {
            "excel_row_number": self.excel_row_number,
            "normalized_data": self.normalized_data,
            "validation_errors": [error.to_dict() for error in self.validation_errors],
        }


@dataclass(frozen=True)
class StudentImportReadResult:
    data_row_count: int
    rows: list[StudentImportRowResult]
    file_errors: list[ImportValidationError]

    @property
    def is_valid(self):
        return not self.file_errors and all(not row.validation_errors for row in self.rows)

    def to_dict(self):
        return {
            "data_row_count": self.data_row_count,
            "rows": [row.to_dict() for row in self.rows],
            "file_errors": [error.to_dict() for error in self.file_errors],
        }


def _error(code, detail, *, field=None, column=None, row_number=None):
    return ImportValidationError(
        code=code,
        detail=detail,
        field=field,
        column=column,
        excel_row_number=row_number,
    )


def _empty_result(error):
    return StudentImportReadResult(data_row_count=0, rows=[], file_errors=[error])


def _validate_xlsx_archive(archive):
    entries = archive.infolist()
    if len(entries) > MAX_XLSX_ZIP_ENTRIES:
        return _error(
            "unsafe_xlsx_archive",
            "يحتوي ملف XLSX على عدد غير طبيعي من الأجزاء المضغوطة.",
        )

    names = [entry.filename for entry in entries]
    if len(names) != len(set(names)):
        return _error(
            "unsafe_xlsx_archive",
            "يحتوي ملف XLSX على أسماء أجزاء مكررة وغير آمنة.",
        )

    total_uncompressed_size = 0
    for entry in entries:
        normalized_name = posixpath.normpath(entry.filename.replace("\\", "/"))
        if (
            entry.flag_bits & 0x1
            or normalized_name.startswith("../")
            or normalized_name.startswith("/")
        ):
            return _error(
                "unsafe_xlsx_archive",
                "يحتوي ملف XLSX على بنية أجزاء غير آمنة أو غير مدعومة.",
            )
        if entry.file_size > MAX_XLSX_PART_UNCOMPRESSED_BYTES:
            return _error(
                "unsafe_xlsx_archive",
                "يحتوي ملف XLSX على جزء غير مضغوط يتجاوز حد الأمان.",
            )
        total_uncompressed_size += entry.file_size
        if total_uncompressed_size > MAX_XLSX_TOTAL_UNCOMPRESSED_BYTES:
            return _error(
                "unsafe_xlsx_archive",
                "يتجاوز الحجم الإجمالي غير المضغوط لملف XLSX حد الأمان.",
            )
        if (
            entry.file_size > 1024 * 1024
            and entry.file_size
            > max(entry.compress_size, 1) * MAX_XLSX_COMPRESSION_RATIO
        ):
            return _error(
                "unsafe_xlsx_archive",
                "تدل نسبة ضغط ملف XLSX على بنية مضغوطة غير آمنة.",
            )
    return None


def _student_sheet_part(archive):
    workbook_root = fromstring(archive.read("xl/workbook.xml"))
    relationships_root = fromstring(
        archive.read("xl/_rels/workbook.xml.rels")
    )
    relationship_targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in relationships_root.findall(
            f"{{{PACKAGE_RELATIONSHIPS_NAMESPACE}}}Relationship"
        )
    }
    for sheet in workbook_root.findall(
        f".//{{{SPREADSHEET_NAMESPACE}}}sheet"
    ):
        if sheet.attrib.get("name") != IMPORT_SHEET_NAME:
            continue
        relationship_id = sheet.attrib.get(
            f"{{{OFFICE_RELATIONSHIPS_NAMESPACE}}}id"
        )
        target = relationship_targets.get(relationship_id)
        if not target:
            return None
        if target.startswith("/"):
            part = target.lstrip("/")
        else:
            part = posixpath.join("xl", target)
        part = posixpath.normpath(part)
        if part.startswith("../") or part not in archive.namelist():
            return None
        return part
    return None


def _xml_cell_has_value(cell):
    formula = cell.find(f"{{{SPREADSHEET_NAMESPACE}}}f")
    if formula is not None:
        return True
    value = cell.find(f"{{{SPREADSHEET_NAMESPACE}}}v")
    if value is not None and value.text not in (None, ""):
        return True
    inline_string = cell.find(f"{{{SPREADSHEET_NAMESPACE}}}is")
    return inline_string is not None and bool(
        "".join(inline_string.itertext()).strip()
    )


def _has_data_after_template_range(archive, sheet_part):
    with archive.open(sheet_part) as sheet_xml:
        for _, element in iterparse(sheet_xml, events=("end",)):
            if element.tag != f"{{{SPREADSHEET_NAMESPACE}}}c":
                continue
            reference = element.attrib.get("r", "")
            row_match = re.search(r"(\d+)$", reference)
            if (
                row_match
                and int(row_match.group(1)) > MAX_TEMPLATE_EXCEL_ROW
                and _xml_cell_has_value(element)
            ):
                return True
            element.clear()
    return False


def _read_upload(source):
    if isinstance(source, (str, Path)):
        path = Path(source)
        size = path.stat().st_size
        if size > MAX_FILE_SIZE_BYTES:
            return None, size
        with path.open("rb") as uploaded_file:
            content = uploaded_file.read(MAX_FILE_SIZE_BYTES + 1)
        return content, len(content)

    size = getattr(source, "size", None)
    if size is not None and size > MAX_FILE_SIZE_BYTES:
        return None, size

    current_position = source.tell() if hasattr(source, "tell") else None
    content = source.read(MAX_FILE_SIZE_BYTES + 1)
    if current_position is not None and hasattr(source, "seek"):
        source.seek(current_position)
    return content, len(content)


def _cell_has_value(cell):
    return cell.value is not None and not (
        isinstance(cell.value, str) and not cell.value.strip()
    )


def _text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _identifier(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return _text(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value).translate(ARABIC_DIGITS_TRANSLATION).replace(" ", "").replace("-", "")


def _phone(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return _text(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value).translate(ARABIC_DIGITS_TRANSLATION)


def _phone_value(values, errors, field, row_number):
    value = _phone(values[field])
    if len(value) > 30:
        _add_field_error(
            errors,
            "max_length",
            "يجب ألا يتجاوز رقم الهاتف 30 محرفًا.",
            field,
            row_number,
        )
    return value


def _add_field_error(errors, code, detail, field, row_number):
    errors.append(
        _error(
            code,
            detail,
            field=field,
            column=HEADER_BY_FIELD[field],
            row_number=row_number,
        )
    )


def _required_text(values, errors, field, row_number, max_length=None):
    value = _text(values[field])
    if not value:
        _add_field_error(errors, "required", "هذا الحقل مطلوب.", field, row_number)
    elif max_length is not None and len(value) > max_length:
        _add_field_error(
            errors,
            "max_length",
            f"يجب ألا يتجاوز هذا الحقل {max_length} محرفًا.",
            field,
            row_number,
        )
    return value


def _optional_text(values, errors, field, row_number, max_length=None):
    value = _text(values[field])
    if max_length is not None and len(value) > max_length:
        _add_field_error(
            errors,
            "max_length",
            f"يجب ألا يتجاوز هذا الحقل {max_length} محرفًا.",
            field,
            row_number,
        )
    return value


def _date_value(value, errors, field, row_number):
    if value is None or (isinstance(value, str) and not value.strip()):
        _add_field_error(errors, "required", "هذا الحقل مطلوب.", field, row_number)
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        candidate = value.strip().translate(ARABIC_DIGITS_TRANSLATION)
        try:
            return date.fromisoformat(candidate).isoformat()
        except ValueError:
            pass
    _add_field_error(
        errors,
        "invalid_date",
        "يجب إدخال تاريخ صريح وصالح بصيغة YYYY-MM-DD.",
        field,
        row_number,
    )
    return None


def _choice(values, errors, field, row_number, choices, *, blank_default=None):
    value = _text(values[field])
    if not value and blank_default is not None:
        return blank_default
    if not value:
        _add_field_error(errors, "required", "هذا الحقل مطلوب.", field, row_number)
        return None
    normalized = choices.get(value)
    if normalized is None:
        _add_field_error(
            errors,
            "invalid_choice",
            "القيمة المدخلة غير معتمدة لهذا الحقل.",
            field,
            row_number,
        )
    return normalized


def _validate_row(cells, row_number):
    values = dict(zip(FIELD_NAMES, (cell.value for cell in cells)))
    errors = []

    for field, cell in zip(FIELD_NAMES, cells):
        if cell.data_type == "f" or (
            isinstance(cell.value, str) and cell.value.lstrip().startswith("=")
        ):
            _add_field_error(
                errors,
                "formula_not_allowed",
                "لا يُسمح باستخدام الصيغ داخل صفوف بيانات الطلاب.",
                field,
                row_number,
            )
            values[field] = None

    student = {
        "first_name": _required_text(values, errors, "first_name", row_number, 100),
        "last_name": _required_text(values, errors, "last_name", row_number, 100),
        "first_name_en": _optional_text(values, errors, "first_name_en", row_number, 100),
        "last_name_en": _optional_text(values, errors, "last_name_en", row_number, 100),
        "father_name": _optional_text(values, errors, "father_name", row_number, 100),
        "mother_name": _optional_text(values, errors, "mother_name", row_number, 100),
        "birth_date": _date_value(values["birth_date"], errors, "birth_date", row_number),
        "gender": _choice(values, errors, "gender", row_number, GENDER_VALUES),
    }

    guardian_fields = (
        "guardian_national_id",
        "guardian_first_name",
        "guardian_last_name",
        "guardian_phone_number",
        "guardian_relationship",
    )
    guardian = None
    if any(_cell_has_value(cells[FIELD_NAMES.index(field)]) for field in guardian_fields):
        national_id = _identifier(values["guardian_national_id"])
        if not national_id:
            _add_field_error(
                errors,
                "required",
                "الرقم الوطني لولي الأمر مطلوب عند إدخال بيانات ولي الأمر.",
                "guardian_national_id",
                row_number,
            )
        elif len(national_id) > 30:
            _add_field_error(errors, "max_length", "يجب ألا يتجاوز الرقم الوطني 30 محرفًا.", "guardian_national_id", row_number)
        elif not national_id.isascii() or not national_id.isdigit():
            _add_field_error(errors, "invalid_identifier", "يجب أن يحتوي الرقم الوطني على أرقام فقط.", "guardian_national_id", row_number)

        guardian = {
            "national_id": national_id,
            "first_name": _required_text(values, errors, "guardian_first_name", row_number, 150),
            "last_name": _required_text(values, errors, "guardian_last_name", row_number, 150),
            "phone_number": _phone_value(
                values, errors, "guardian_phone_number", row_number
            ),
            "relationship": _required_text(values, errors, "guardian_relationship", row_number, 50),
        }

    blood_type = _optional_text(values, errors, "blood_type", row_number, 10).upper()
    if blood_type and blood_type not in BLOOD_TYPES:
        _add_field_error(errors, "invalid_blood_type", "زمرة الدم المدخلة غير معتمدة.", "blood_type", row_number)

    health_profile = {
        "blood_type": blood_type,
        "chronic_diseases": _optional_text(values, errors, "chronic_diseases", row_number),
        "allergies": _optional_text(values, errors, "allergies", row_number),
        "permanent_medications": _optional_text(values, errors, "permanent_medications", row_number),
        "special_health_needs": _optional_text(values, errors, "special_health_needs", row_number),
        "emergency_contact_name": _optional_text(values, errors, "emergency_contact_name", row_number, 150),
        "emergency_contact_phone": _phone_value(
            values, errors, "emergency_contact_phone", row_number
        ),
        "health_notes": _optional_text(values, errors, "health_notes", row_number),
    }

    academic_year = _required_text(values, errors, "academic_year", row_number)
    year_match = ACADEMIC_YEAR_PATTERN.fullmatch(academic_year.translate(ARABIC_DIGITS_TRANSLATION))
    if academic_year and (
        year_match is None or int(year_match.group(2)) != int(year_match.group(1)) + 1
    ):
        _add_field_error(errors, "invalid_academic_year", "يجب أن تكون السنة الدراسية بصيغة YYYY/YYYY لسنتين متتاليتين.", "academic_year", row_number)
    elif year_match:
        academic_year = f"{year_match.group(1)}/{year_match.group(2)}"

    enrollment = {
        "academic_year": academic_year,
        "stage": _choice(values, errors, "stage", row_number, STAGE_VALUES),
        "grade_level": _required_text(values, errors, "grade_level", row_number, 100),
        "section": _required_text(values, errors, "section", row_number, 80),
        "enrollment_date": _date_value(values["enrollment_date"], errors, "enrollment_date", row_number),
        "usual_arrival_method": _choice(values, errors, "usual_arrival_method", row_number, TRANSPORTATION_VALUES, blank_default="school_bus"),
        "usual_departure_method": _choice(values, errors, "usual_departure_method", row_number, TRANSPORTATION_VALUES, blank_default="school_bus"),
    }

    return StudentImportRowResult(
        excel_row_number=row_number,
        normalized_data={
            "student": student,
            "guardian": guardian,
            "health_profile": health_profile,
            "enrollment": enrollment,
        },
        validation_errors=errors,
    )


def read_student_import_xlsx(source, *, filename=None):
    effective_filename = filename or getattr(source, "name", None)
    if not effective_filename or Path(str(effective_filename)).suffix.lower() != ".xlsx":
        return _empty_result(
            _error("unsupported_extension", "يجب أن يكون الملف بصيغة XLSX.")
        )

    try:
        content, size = _read_upload(source)
    except (OSError, AttributeError, TypeError):
        return _empty_result(_error("file_read_error", "تعذرت قراءة الملف المرفوع."))

    if size > MAX_FILE_SIZE_BYTES:
        return _empty_result(
            _error("file_too_large", f"يجب ألا يتجاوز حجم الملف {MAX_FILE_SIZE_BYTES // (1024 * 1024)} ميغابايت.")
        )
    if not content:
        return _empty_result(_error("empty_file", "الملف فارغ."))
    if not isinstance(content, (bytes, bytearray)):
        return _empty_result(_error("file_read_error", "تعذرت قراءة الملف المرفوع."))

    stream = BytesIO(content)
    if not zipfile.is_zipfile(stream):
        return _empty_result(_error("invalid_xlsx", "محتوى الملف ليس مصنف XLSX صالحًا."))

    try:
        with zipfile.ZipFile(stream) as archive:
            archive_error = _validate_xlsx_archive(archive)
            if archive_error:
                return _empty_result(archive_error)

            required_parts = {
                "[Content_Types].xml",
                "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels",
            }
            if not required_parts.issubset(archive.namelist()):
                return _empty_result(_error("invalid_xlsx", "بنية ملف XLSX غير صالحة."))

            sheet_part = _student_sheet_part(archive)
            if sheet_part and _has_data_after_template_range(archive, sheet_part):
                return _empty_result(
                    _error(
                        "data_outside_template_range",
                        "توجد بيانات فعلية بعد الصف 1501 خارج نطاق قالب الاستيراد المعتمد.",
                    )
                )

        stream.seek(0)
        workbook = load_workbook(
            stream,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (
        OSError,
        AttributeError,
        ValueError,
        KeyError,
        IndexError,
        EOFError,
        OverflowError,
        ParseError,
        InvalidFileException,
        zipfile.LargeZipFile,
        RuntimeError,
        TypeError,
        zipfile.BadZipFile,
    ):
        return _empty_result(_error("invalid_xlsx", "تعذر فتح ملف XLSX لأن بنيته تالفة أو غير مدعومة."))

    try:
        if IMPORT_SHEET_NAME not in workbook.sheetnames:
            return _empty_result(
                _error("missing_sheet", f'ورقة العمل المطلوبة "{IMPORT_SHEET_NAME}" غير موجودة.')
            )

        sheet = workbook[IMPORT_SHEET_NAME]
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        actual_headers = tuple(_text(cell.value) for cell in header_cells[: len(EXPECTED_HEADERS)])
        has_extra_headers = any(_cell_has_value(cell) for cell in header_cells[len(EXPECTED_HEADERS) :])
        if actual_headers != EXPECTED_HEADERS or has_extra_headers:
            return _empty_result(
                _error("invalid_headers", "عناوين الأعمدة أو ترتيبها لا يطابق قالب الاستيراد المعتمد.")
            )

        rows = []
        data_row_count = 0
        for excel_row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                max_row=MAX_TEMPLATE_EXCEL_ROW,
                max_col=len(EXPECTED_HEADERS),
            ),
            start=2,
        ):
            if not any(_cell_has_value(cell) for cell in row):
                continue
            data_row_count += 1
            rows.append(_validate_row(row, excel_row_number))
        if not rows:
            return _empty_result(_error("no_data_rows", "لا يحتوي الملف على صفوف بيانات طلاب."))

        return StudentImportReadResult(
            data_row_count=data_row_count,
            rows=rows,
            file_errors=[],
        )
    finally:
        workbook.close()
